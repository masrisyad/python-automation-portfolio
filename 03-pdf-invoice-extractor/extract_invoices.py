"""
PDF Invoice Data Extractor
==========================
Reads every invoice PDF in a folder and extracts header fields and line
items into a structured Excel workbook, with built-in validation.

Replaces manual data entry from hundreds of supplier invoices into
accounting software or spreadsheets.

Output workbook
- Invoices    : one row per invoice (vendor, number, dates, bill-to, totals)
- Line Items  : one row per line, linked by invoice number
- Issues      : invoices where the extracted numbers do not add up, or a
                field could not be found (for manual review)

Usage
    python extract_invoices.py                    # invoices/*.pdf -> output/invoices.xlsx
    python extract_invoices.py --input scans --output result.xlsx
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
import pdfplumber
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("extract")

DATE = r"(\d{4}-\d{2}-\d{2}|\d{1,2}/\d{1,2}/\d{4}|\w+ \d{1,2}, \d{4})"
PATTERNS = {
    "invoice_number": re.compile(r"Invoice\s*(?:Number|No\.?|#)\s*[:#]?\s*([A-Z0-9][A-Z0-9-]+)", re.I),
    "invoice_date": re.compile(r"Invoice\s*Date\s*:?\s*" + DATE, re.I),
    "due_date": re.compile(r"Due\s*Date\s*:?\s*" + DATE, re.I),
    "bill_to": re.compile(r"Bill\s*To\s*:?\s*(.+)", re.I),
    "subtotal": re.compile(r"Subtotal\s*:?\s*\$?\s*([\d,]+\.\d{2})", re.I),
    "tax": re.compile(r"Tax\s*(?:\([^)]*\))?\s*:?\s*\$?\s*([\d,]+\.\d{2})", re.I),
    "total": re.compile(r"(?<![A-Za-z])Total\s*(?:Due|Amount)?\s*:?\s*\$?\s*([\d,]+\.\d{2})", re.I),
}


@dataclass
class Invoice:
    file: str
    vendor: str | None = None
    invoice_number: str | None = None
    invoice_date: str | None = None
    due_date: str | None = None
    bill_to: str | None = None
    subtotal: float | None = None
    tax: float | None = None
    total: float | None = None
    line_items: list[dict] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)


def money(text: str | None) -> float | None:
    if text is None:
        return None
    return float(text.replace(",", "").lstrip("$"))


def parse_line_items(tables: list[list[list[str | None]]], inv_no: str | None) -> list[dict]:
    items = []
    for table in tables:
        if not table or not table[0]:
            continue
        header = [(c or "").strip().lower() for c in table[0]]
        if "description" not in header or not any("amount" in h or "total" in h for h in header):
            continue
        for row in table[1:]:
            cells = [(c or "").strip() for c in row]
            if len(cells) < 4 or not cells[0]:
                continue
            try:
                items.append({
                    "invoice_number": inv_no,
                    "description": cells[0],
                    "quantity": float(cells[1]),
                    "unit_price": money(cells[2]),
                    "amount": money(cells[3]),
                })
            except ValueError:
                continue
    return items


def extract_one(path: Path) -> Invoice:
    inv = Invoice(file=path.name)
    with pdfplumber.open(path) as pdf:
        text = "\n".join((page.extract_text() or "") for page in pdf.pages)
        tables = [t for page in pdf.pages for t in page.extract_tables()]

    lines = [l.strip() for l in text.splitlines() if l.strip()]
    inv.vendor = lines[0] if lines else None

    found = {}
    for key, pattern in PATTERNS.items():
        m = pattern.search(text)
        found[key] = m.group(1).strip() if m else None
    inv.invoice_number = found["invoice_number"]
    inv.invoice_date = found["invoice_date"]
    inv.due_date = found["due_date"]
    inv.bill_to = found["bill_to"]
    inv.subtotal = money(found["subtotal"])
    inv.tax = money(found["tax"])
    inv.total = money(found["total"])
    inv.line_items = parse_line_items(tables, inv.invoice_number)

    # validation
    for k in ("invoice_number", "invoice_date", "total"):
        if getattr(inv, k) is None:
            inv.issues.append(f"missing {k}")
    if not inv.line_items:
        inv.issues.append("no line items found")
    elif inv.subtotal is not None:
        calc = round(sum(i["amount"] or 0 for i in inv.line_items), 2)
        if abs(calc - inv.subtotal) > 0.01:
            inv.issues.append(f"line items sum {calc:.2f} != subtotal {inv.subtotal:.2f}")
    if None not in (inv.subtotal, inv.tax, inv.total) and abs(inv.subtotal + inv.tax - inv.total) > 0.01:
        inv.issues.append("subtotal + tax != total")
    return inv


def style(ws, money_cols: set[str]) -> None:
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for idx, cell in enumerate(ws[1], start=1):
        col = get_column_letter(idx)
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in ws[col])
        ws.column_dimensions[col].width = min(max(12, max_len + 2), 50)
        if str(cell.value) in money_cols:
            for c in ws[col][1:]:
                c.number_format = '"$"#,##0.00'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_output(invoices: list[Invoice], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    header_rows = [{
        "file": i.file, "vendor": i.vendor, "invoice_number": i.invoice_number,
        "invoice_date": i.invoice_date, "due_date": i.due_date, "bill_to": i.bill_to,
        "subtotal": i.subtotal, "tax": i.tax, "total": i.total,
        "line_items": len(i.line_items), "status": "OK" if not i.issues else "REVIEW",
    } for i in invoices]
    line_rows = [li | {"file": i.file} for i in invoices for li in i.line_items]
    issue_rows = [{"file": i.file, "invoice_number": i.invoice_number, "issue": msg}
                  for i in invoices for msg in i.issues]

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(header_rows).to_excel(writer, sheet_name="Invoices", index=False)
        pd.DataFrame(line_rows or [{"invoice_number": None}]).to_excel(writer, sheet_name="Line Items", index=False)
        pd.DataFrame(issue_rows or [{"file": None, "invoice_number": None, "issue": "No issues found"}]).to_excel(writer, sheet_name="Issues", index=False)
        style(writer.book["Invoices"], {"subtotal", "tax", "total"})
        style(writer.book["Line Items"], {"unit_price", "amount"})
        style(writer.book["Issues"], set())
    log.info("Wrote %s", path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", default="invoices", help="folder containing PDF invoices")
    parser.add_argument("--output", default="output/invoices.xlsx")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    files = sorted(Path(args.input).glob("*.pdf"))
    if not files:
        log.error("No PDF files in %s", args.input)
        return 1
    invoices = []
    for f in files:
        inv = extract_one(f)
        invoices.append(inv)
        log.info("%s -> %s | %s | total %s | %d items | %s",
                 f.name, inv.invoice_number, inv.vendor, inv.total, len(inv.line_items),
                 "OK" if not inv.issues else "; ".join(inv.issues))
    write_output(invoices, Path(args.output))
    ok = sum(1 for i in invoices if not i.issues)
    log.info("Done: %d invoices, %d OK, %d need review.", len(invoices), ok, len(invoices) - ok)
    return 0


if __name__ == "__main__":
    sys.exit(main())
