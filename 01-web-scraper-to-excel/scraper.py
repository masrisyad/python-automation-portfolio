"""
Web Scraper -> CSV / Excel
==========================
Scrapes the full product catalogue of https://books.toscrape.com (a public
sandbox site built for scraping practice) and exports clean, analysis-ready
data to CSV and a formatted Excel workbook.

Features
- Handles pagination automatically (stops when the site returns 404)
- Retries with exponential backoff on network errors
- Polite rate limiting between requests
- Clean typed fields (price as float, rating as integer, stock as boolean)
- Excel output with bold headers, auto-filter, frozen header row,
  currency formatting and sensible column widths

Usage
    python scraper.py                      # scrape everything, write CSV + XLSX
    python scraper.py --max-pages 3        # quick test run
    python scraper.py --format csv --out output/books
"""

from __future__ import annotations

import argparse
import csv
import logging
import sys
import time
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/catalogue/page-{page}.html"
RATING_MAP = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
HEADERS = {"User-Agent": "Mozilla/5.0 (portfolio scraper; contact via Upwork)"}

log = logging.getLogger("scraper")


@dataclass
class Book:
    title: str
    price_gbp: float
    rating: int
    in_stock: bool
    url: str


def fetch(session: requests.Session, url: str, retries: int = 3, backoff: float = 1.5) -> requests.Response | None:
    """GET a URL with retries. Returns None on 404 (end of catalogue)."""
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, headers=HEADERS, timeout=20)
            if resp.status_code == 404:
                return None
            resp.raise_for_status()
            return resp
        except requests.RequestException as exc:
            wait = backoff ** attempt
            log.warning("Attempt %d/%d failed for %s: %s (retry in %.1fs)", attempt, retries, url, exc, wait)
            time.sleep(wait)
    raise RuntimeError(f"Giving up on {url} after {retries} attempts")


def parse_page(html: str, page_url: str) -> list[Book]:
    soup = BeautifulSoup(html, "html.parser")
    books: list[Book] = []
    for card in soup.select("article.product_pod"):
        link = card.select_one("h3 a")
        price_text = card.select_one("p.price_color").get_text(strip=True)
        rating_word = card.select_one("p.star-rating")["class"][1]
        stock_text = card.select_one("p.instock.availability").get_text(strip=True)
        digits = "".join(ch for ch in price_text if ch.isdigit() or ch == ".")
        books.append(
            Book(
                title=link["title"].strip(),
                price_gbp=float(digits),
                rating=RATING_MAP.get(rating_word, 0),
                in_stock="In stock" in stock_text,
                url=urljoin(page_url, link["href"]),
            )
        )
    return books


def scrape(max_pages: int | None, delay: float) -> list[Book]:
    session = requests.Session()
    all_books: list[Book] = []
    page = 1
    while True:
        if max_pages and page > max_pages:
            break
        url = BASE_URL.format(page=page)
        resp = fetch(session, url)
        if resp is None:
            log.info("Page %d not found, catalogue finished.", page)
            break
        books = parse_page(resp.text, url)
        if not books:
            break
        all_books.extend(books)
        log.info("Page %d: %d books (total %d)", page, len(books), len(all_books))
        page += 1
        time.sleep(delay)
    return all_books


def save_csv(books: list[Book], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=[f.name for f in fields(Book)])
        writer.writeheader()
        writer.writerows(asdict(b) for b in books)
    log.info("Wrote %s", path)


def save_xlsx(books: list[Book], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"

    headers = ["Title", "Price (GBP)", "Rating (1-5)", "In Stock", "URL"]
    ws.append(headers)
    for b in books:
        ws.append([b.title, b.price_gbp, b.rating, "Yes" if b.in_stock else "No", b.url])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = '"GBP "#,##0.00'

    widths = {1: 60, 2: 13, 3: 13, 4: 10, 5: 70}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
    log.info("Wrote %s", path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--max-pages", type=int, default=None, help="limit number of pages (default: all)")
    parser.add_argument("--delay", type=float, default=0.5, help="seconds to wait between requests")
    parser.add_argument("--out", default="output/books", help="output path without extension")
    parser.add_argument("--format", choices=["csv", "xlsx", "both"], default="both")
    args = parser.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    books = scrape(args.max_pages, args.delay)
    if not books:
        log.error("No data scraped.")
        return 1

    out = Path(args.out)
    if args.format in ("csv", "both"):
        save_csv(books, out.with_suffix(".csv"))
    if args.format in ("xlsx", "both"):
        save_xlsx(books, out.with_suffix(".xlsx"))

    log.info("Done: %d books exported.", len(books))
    return 0


if __name__ == "__main__":
    sys.exit(main())
