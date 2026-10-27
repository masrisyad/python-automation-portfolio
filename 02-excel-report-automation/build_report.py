"""
Excel Report Automation
=======================
Merges every CSV export in an input folder, cleans the data, and produces a
management-ready Excel workbook with KPIs, pivot tables and native charts.

Replaces a manual monthly process (open 6 files, copy-paste, fix typos,
build pivots, make charts) with a single command.

Usage
    python build_report.py                       # input/*.csv -> output/sales_report.xlsx
    python build_report.py --input data --output report.xlsx
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("report")

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


# --------------------------------------------------------------------------- #
# 1. Load + clean
# --------------------------------------------------------------------------- #
def load_inputs(folder: Path) -> pd.DataFrame:
    files = sorted(folder.glob("*.csv"))
    if not files:
        raise FileNotFoundError(f"No CSV files found in {folder}")
    frames = []
    for f in files:
        df = pd.read_csv(f, dtype=str)
        df["source_file"] = f.name
        frames.append(df)
        log.info("Loaded %s (%d rows)", f.name, len(df))
    return pd.concat(frames, ignore_index=True)


def clean(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    stats = {"rows_in": len(df)}

    # normalise text columns: trim and collapse spaces, then merge spelling
    # variants ("SOUTH", " south ") into the most common spelling ("South")
    def canonical(series: pd.Series) -> pd.Series:
        tidy = series.astype(str).str.split().str.join(" ")
        key = tidy.str.lower()
        best = tidy.groupby(key).agg(lambda x: x.mode().iloc[0])
        return key.map(best)

    for col in ("customer", "region", "product"):
        df[col] = canonical(df[col])

    # types
    df["order_date"] = pd.to_datetime(df["order_date"], format="mixed", errors="coerce")
    df["quantity"] = pd.to_numeric(df["quantity"], errors="coerce")
    df["unit_price"] = pd.to_numeric(df["unit_price"], errors="coerce")

    # drop exact duplicates (same order id + same content)
    before = len(df)
    df = df.drop_duplicates(subset=[c for c in df.columns if c != "source_file"])
    stats["duplicates_removed"] = before - len(df)

    # drop rows that cannot be used
    before = len(df)
    df = df.dropna(subset=["order_date", "quantity", "unit_price"])
    stats["invalid_rows_removed"] = before - len(df)

    df["revenue"] = (df["quantity"] * df["unit_price"]).round(2)
    df["month"] = df["order_date"].dt.to_period("M").astype(str)
    df = df.sort_values("order_date").reset_index(drop=True)
    stats["rows_out"] = len(df)
    return df, stats


# --------------------------------------------------------------------------- #
# 2. Aggregate
# --------------------------------------------------------------------------- #
def aggregate(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    monthly = (df.groupby("month")
                 .agg(orders=("order_id", "count"), units=("quantity", "sum"), revenue=("revenue", "sum"))
                 .reset_index())
    monthly["revenue_growth_pct"] = (monthly["revenue"].pct_change() * 100).round(1)

    by_region = (df.groupby("region")
                   .agg(orders=("order_id", "count"), revenue=("revenue", "sum"))
                   .sort_values("revenue", ascending=False).reset_index())
    by_region["share_pct"] = (by_region["revenue"] / by_region["revenue"].sum() * 100).round(1)

    by_product = (df.groupby("product")
                    .agg(units=("quantity", "sum"), revenue=("revenue", "sum"),
                         avg_price=("unit_price", "mean"))
                    .sort_values("revenue", ascending=False).reset_index())
    by_product["avg_price"] = by_product["avg_price"].round(2)

    top_customers = (df.groupby("customer")
                       .agg(orders=("order_id", "count"), revenue=("revenue", "sum"))
                       .sort_values("revenue", ascending=False).head(10).reset_index())

    region_month = df.pivot_table(index="region", columns="month", values="revenue", aggfunc="sum", fill_value=0).round(2)
    region_month = region_month.reset_index()

    return {"Monthly": monthly, "By Region": by_region, "By Product": by_product,
            "Top Customers": top_customers, "Region x Month": region_month}


# --------------------------------------------------------------------------- #
# 3. Write Excel
# --------------------------------------------------------------------------- #
def style_sheet(ws, money_cols: set[str] | None = None, pct_cols: set[str] | None = None) -> None:
    money_cols = money_cols or set()
    pct_cols = pct_cols or set()
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for idx, cell in enumerate(ws[1], start=1):
        header = str(cell.value)
        col = get_column_letter(idx)
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in ws[col])
        ws.column_dimensions[col].width = min(max(12, max_len + 2), 40)
        if header in money_cols or header.startswith("2026-"):
            for c in ws[col][1:]:
                c.number_format = '"$"#,##0.00'
        elif header in pct_cols:
            for c in ws[col][1:]:
                c.number_format = '0.0"%"'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_summary(wb, df: pd.DataFrame, stats: dict, tables: dict) -> None:
    ws = wb.create_sheet("Summary", 0)
    ws["A1"] = "Sales Report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Period: {df['order_date'].min():%d %b %Y} to {df['order_date'].max():%d %b %Y}"

    kpis = [
        ("Total revenue", df["revenue"].sum(), '"$"#,##0.00'),
        ("Total orders", len(df), "#,##0"),
        ("Units sold", int(df["quantity"].sum()), "#,##0"),
        ("Average order value", df["revenue"].mean(), '"$"#,##0.00'),
        ("Best month", tables["Monthly"].loc[tables["Monthly"]["revenue"].idxmax(), "month"], "@"),
        ("Top region", tables["By Region"].iloc[0]["region"], "@"),
        ("Top product", tables["By Product"].iloc[0]["product"], "@"),
    ]
    ws["A4"] = "KPI"
    ws["B4"] = "Value"
    for c in ("A4", "B4"):
        ws[c].fill = HEADER_FILL
        ws[c].font = HEADER_FONT
    for i, (label, value, fmt) in enumerate(kpis, start=5):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = value
        ws[f"B{i}"].number_format = fmt

    r = 5 + len(kpis) + 1
    ws[f"A{r}"] = "Data quality"
    ws[f"B{r}"] = "Count"
    for c in (f"A{r}", f"B{r}"):
        ws[c].fill = HEADER_FILL
        ws[c].font = HEADER_FONT
    for i, (k, v) in enumerate(stats.items(), start=r + 1):
        ws[f"A{i}"] = k.replace("_", " ").capitalize()
        ws[f"B{i}"] = v
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 18


def add_charts(wb, tables: dict) -> None:
    monthly = wb["Monthly"]
    n = len(tables["Monthly"]) + 1
    line = LineChart()
    line.title = "Monthly revenue"
    line.y_axis.title = "USD"
    line.x_axis.title = "Month"
    line.add_data(Reference(monthly, min_col=4, min_row=1, max_row=n), titles_from_data=True)
    line.set_categories(Reference(monthly, min_col=1, min_row=2, max_row=n))
    line.height, line.width = 8, 16
    monthly.add_chart(line, "G2")

    region = wb["By Region"]
    n = len(tables["By Region"]) + 1
    bar = BarChart()
    bar.title = "Revenue by region"
    bar.add_data(Reference(region, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    bar.set_categories(Reference(region, min_col=1, min_row=2, max_row=n))
    bar.height, bar.width = 8, 14
    region.add_chart(bar, "F2")

    product = wb["By Product"]
    n = len(tables["By Product"]) + 1
    bar2 = BarChart()
    bar2.type = "bar"
    bar2.title = "Revenue by product"
    bar2.add_data(Reference(product, min_col=3, min_row=1, max_row=n), titles_from_data=True)
    bar2.set_categories(Reference(product, min_col=1, min_row=2, max_row=n))
    bar2.height, bar2.width = 8, 14
    product.add_chart(bar2, "G2")

    rm = wb["Region x Month"]
    last_col = get_column_letter(rm.max_column)
    rm.conditional_formatting.add(f"B2:{last_col}{rm.max_row}",
                                  ColorScaleRule(start_type="min", start_color="F8696B",
                                                 mid_type="percentile", mid_value=50, mid_color="FFEB84",
                                                 end_type="max", end_color="63BE7B"))


def write_workbook(df: pd.DataFrame, stats: dict, tables: dict, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    clean_cols = ["order_id", "order_date", "customer", "region", "product", "quantity", "unit_price", "revenue", "source_file"]
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
        for name, table in tables.items():
            table.to_excel(writer, sheet_name=name, index=False)
        df[clean_cols].to_excel(writer, sheet_name="Clean Data", index=False)

        wb = writer.book
        style_sheet(wb["Monthly"], money_cols={"revenue"}, pct_cols={"revenue_growth_pct"})
        style_sheet(wb["By Region"], money_cols={"revenue"}, pct_cols={"share_pct"})
        style_sheet(wb["By Product"], money_cols={"revenue", "avg_price"})
        style_sheet(wb["Top Customers"], money_cols={"revenue"})
        style_sheet(wb["Region x Month"])
        style_sheet(wb["Clean Data"], money_cols={"unit_price", "revenue"})
        add_charts(wb, tables)
        write_summary(wb, df, stats, tables)
    log.info("Wrote %s", path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", default="input", help="folder with CSV exports")
    parser.add_argument("--output", default="output/sales_report.xlsx")
    args = parser.parse_args(argv)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s", datefmt="%H:%M:%S")

    raw = load_inputs(Path(args.input))
    df, stats = clean(raw)
    log.info("Clean: %s", stats)
    tables = aggregate(df)
    write_workbook(df, stats, tables, Path(args.output))
    return 0


if __name__ == "__main__":
    sys.exit(main())
